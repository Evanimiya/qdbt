# -*- coding: utf-8 -*-
"""견고성 테스트 공용 유틸 (QDBT_견고성테스트_20260711).

'엉망으로 접수된 견적서' 시나리오를 openpyxl 합성 워크북으로 만들고,
extract_by_mapping / stitch / 통화정합 / 집계 를 함수 직접 호출로 검증한다.
포그라운드 블로킹 프로세스(정적 서버 등)는 절대 쓰지 않는다.

각 시나리오 스크립트는 이 모듈의 Recorder에 결과를 append(jsonl)하고,
run_all.py 가 이를 모아 docs/QDBT_견고성테스트_20260711.md 매트릭스로 렌더한다.
"""
import os
import sys
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

# 산출물 경로
FIX_DIR = ROOT / "tests" / "rob" / "_fixtures"
RESULTS = ROOT / "tests" / "rob" / "_results.jsonl"
FIX_DIR.mkdir(parents=True, exist_ok=True)


# ── 워크북 빌더 ──────────────────────────────────
def build_wb(name, sheets):
    """합성 워크북 생성.

    sheets: [(sheet_name, rows_2d, merges)] 형태.
      rows_2d: [[셀, 셀, ...], ...]  (1행=엑셀 1행, None 허용)
      merges:  ["A1:A3", ...]  좌상단 값이 병합영역 대표값.
    반환: 저장된 파일 경로(str).
    """
    from openpyxl import Workbook
    wb = Workbook()
    first = True
    for sname, rows, merges in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = sname
        first = False
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row, start=1):
                if val is not None:
                    ws.cell(row=r, column=c, value=val)
        for m in (merges or []):
            ws.merge_cells(m)
    path = FIX_DIR / name
    wb.save(path)
    wb.close()
    return str(path)


# ── 집계 헬퍼 (queries.recompute_subtotal 미러) ──
def leaf_total(items):
    """정본 잎 합계 = Σ amount (merge_status·category_header 제외).

    queries.recompute_subtotal 은 get_items(headers=False) = is_header=0 항목의
    amount 합. insert_items_bulk 은 (is_category_header or merge_status) → is_header=1.
    따라서 이 식이 저장 후 집계 총액과 동일하다(통화정합은 amount에 이미 반영).
    """
    tot = 0.0
    for it in items:
        if it.get("merge_status") or it.get("is_category_header") or it.get("bom_part"):
            continue   # 부품(bom_part)은 단가 내역 → 총액 제외(is_header=1과 동일)
        amt = it.get("amount")
        if amt:
            tot += amt
    return round(tot, 4)


def krw_leaf_total_from_extract(res):
    """extract_by_mapping 결과(dict)에서 정본 잎 KRW 총액."""
    return leaf_total(res["items"])


# ── 결과 기록 ────────────────────────────────────
def reset_results():
    if RESULTS.exists():
        RESULTS.unlink()


class Recorder:
    """시나리오 결과를 jsonl 로 누적. 한 스크립트가 여러 케이스 기록 가능."""

    def __init__(self, group):
        self.group = group
        self.rows = []

    def add(self, cid, title, inp, expected, actual, verdict, root=""):
        """verdict: 'PASS' | 'FAIL' | 'WARN'."""
        rec = {
            "group": self.group, "id": cid, "title": title,
            "input": inp, "expected": expected, "actual": actual,
            "verdict": verdict, "root": root,
        }
        self.rows.append(rec)
        mark = {"PASS": "✅", "FAIL": "❌", "WARN": "⚠️"}.get(verdict, "?")
        print(f"  {mark} [{cid}] {title} :: {verdict}")
        if verdict != "PASS":
            print(f"      기대: {expected}")
            print(f"      실제: {actual}")
            if root:
                print(f"      근본원인: {root}")
        return rec

    def flush(self):
        with open(RESULTS, "a", encoding="utf-8") as f:
            for r in self.rows:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        n = len(self.rows)
        npass = sum(1 for r in self.rows if r["verdict"] == "PASS")
        print(f"[{self.group}] {npass}/{n} PASS  → 기록 {RESULTS.name}")
        return self.rows


def approx(a, b, tol=1.0):
    return abs((a or 0) - (b or 0)) <= tol
