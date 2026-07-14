# -*- coding: utf-8 -*-
"""시나리오 14: 재구성층(#2) · 드래그 재레벨링(#4) · 레벨별 엑셀(#3).

비파괴 오버레이(map_config.recon): 신규 분류·항목/묶음 이동·재레벨. 정본 불변·총액 Δ=0.
"""
import io
import os
import contextlib
import _util
from _util import build_wb, approx, Recorder


def _setup(rows, mp):
    tmp = os.path.join(_util.FIX_DIR, "s14.db")
    if os.path.exists(tmp):
        os.remove(tmp)
    with contextlib.redirect_stdout(io.StringIO()):
        import importlib
        import db.schema as schema
        import db.queries as q
        importlib.reload(schema); importlib.reload(q)
        q.DB_PATH = tmp; schema.DB_PATH = tmp
        schema.init_db(tmp, reset=True); schema.migrate_db(tmp)
        from extractors.stitch import stitch_sheets
        p = build_wb("s14.xlsx", [("S", rows, [])])
        sres = stitch_sheets(p, [{"sheet": "S", "mapping": mp, "header_row": 1}])
        pid = q.create_project("P"); bid = q.create_bid(pid, "B")
        sid = q.create_submission(bid, "V", "c.xlsx", "/t", "xlsx")
        q.insert_items_bulk(sid, sres["items"]); q.update_submission(sid, extraction_status="done")
        q.recompute_subtotal(sid)
    return q, sid


def _tree_paths(q, sid, view="recon"):
    t = q.build_items_tree(sid, view=view)
    out = []
    def w(ns):
        for n in ns:
            out.append((n["path"], n["level"]))
            w(n["children"])
    w(t["tree"])
    return out, t["total"]


def run():
    rec = Recorder("S14 재구성층")
    ROWS = [["대분류", "중분류", "품명", "금액"],
            ["재료비", "기구부", "볼트", 1000],
            ["재료비", "전장부", "너트", 500],
            ["노무비", "설치", "설치공", 2000]]
    MP = {1: "cat1", 2: "cat2", 3: "name", 4: "amount"}

    # ── 14-1. 신규 분류 + 오버레이 이동 → Δ=0, 정본 불변, 뷰 전환 ──
    q, sid = _setup(ROWS, MP)
    _, t0 = _tree_paths(q, sid)
    recon = q.get_recon(sid)
    volt = [dict(i)["item_id"] for i in q.get_items(sid) if dict(i)["name_normalized"] == "볼트"][0]
    q.save_recon(sid, {"nodes": [{"id": "u1", "name": "신규중", "level": 2, "parent_path": "재료비"}],
                       "moves": {volt: "재료비 > 신규중"}})
    paths_r, t1 = _tree_paths(q, sid, "recon")
    paths_o, t2 = _tree_paths(q, sid, "original")
    # 정본 path 불변 확인
    orig_volt_path = [dict(i)["path"] for i in q.get_items(sid) if dict(i)["name_normalized"] == "볼트"][0]
    moved = any(p == "재료비 > 신규중 > 볼트" for p, _ in paths_r)
    ok = (moved and approx(t0, 3500) and approx(t1, 3500) and approx(t2, 3500)
          and orig_volt_path == "재료비 > 기구부"
          and any(p == "재료비 > 기구부 > 볼트" for p, _ in paths_o))
    rec.add("14-1", "신규 분류 + 오버레이 이동 (#2)",
            "신규중(레벨2) 생성 + 볼트 이동",
            "재구성뷰=볼트 신규중 아래, original뷰=원래대로, 정본 path 불변, 총액 3500 Δ=0",
            f"이동={moved} 총액 recon/orig={t1}/{t2} 정본볼트={orig_volt_path}",
            "PASS" if ok else "FAIL",
            "" if ok else "오버레이 이동/뷰전환/Δ=0 오류")

    # ── 14-2. 재레벨링(레벨만·서브트리 동반·클램프·역전 방지) (#4) ──
    q, sid = _setup(ROWS, MP)
    from web.blueprints.submissions import _recon_relevel
    r1 = _recon_relevel(sid, "재료비 > 기구부", 3)      # 중→소, 볼트 5→6 동반
    lv = dict(_tree_paths(q, sid)[0])
    _, tot = _tree_paths(q, sid)
    r2 = _recon_relevel(sid, "재료비 > 기구부", 1)      # 역전 시도 → 부모+1로 클램프
    r3 = _recon_relevel(sid, "재료비 > 기구부", 6)      # 서브트리 6 초과 → 클램프
    ok = (r1.get("level") == 3 and lv.get("재료비 > 기구부") == 3
          and lv.get("재료비 > 기구부 > 볼트") == 6
          and r2.get("level") == 2               # 역전 방지(부모 재료비=1 → 최소 2)
          and r3.get("ok") and approx(tot, 3500))
    rec.add("14-2", "드래그 재레벨링 (#4)",
            "기구부 2→3(볼트 동반) · 역전(→1) · 과도(→6) 클램프",
            "기구부=3·볼트=6, 역전 클램프=2, 과도 클램프(≤6), 총액 Δ=0",
            f"→3={r1.get('level')} 볼트={lv.get('재료비 > 기구부 > 볼트')} 역전→{r2.get('level')} 총액={tot}",
            "PASS" if ok else "FAIL",
            "" if ok else "재레벨/클램프/역전방지 오류")

    # ── 14-3. 레벨별 엑셀 (#3): 열·부품 제외·합계 ──
    q, sid = _setup([
        ["대분류", "품명", "수량", "단가", "금액", "부품", "부품수량", "부품단가", "파트금액"],
        ["재료비", "납블록", 2, 300, 600, None, None, None, None],
        [None, None, None, None, None, "순납강판", 3, 50, 150],
        [None, None, None, None, None, "볼트", 1, 150, 150],
    ], None)  # BOM은 extract 경로 — 아래서 직접 추출
    # BOM은 stitch가 아니라 extract 경로 → 재추출·재삽입
    import _util as U
    from extractors.extract_by_mapping import extract_by_mapping
    with contextlib.redirect_stdout(io.StringIO()):
        q.execute if False else None
        # 재삽입
        with q.get_conn() as c:
            c.execute("DELETE FROM submission_items WHERE submission_id=?", (sid,))
        p = U.build_wb("s14b.xlsx", [("S", [
            ["대분류", "품명", "수량", "단가", "금액", "부품", "부품수량", "부품단가", "파트금액"],
            ["재료비", "납블록", 2, 300, 600, None, None, None, None],
            [None, None, None, None, None, "순납강판", 3, 50, 150],
            [None, None, None, None, None, "볼트", 1, 150, 150],
        ], [])])
        res = extract_by_mapping(p, "S", {1: "cat1", 2: "name", 3: "qty", 4: "price", 5: "amount",
                                          6: "part", 7: "part_qty", 8: "part_price", 9: "part_amount"}, 1)
        q.insert_items_bulk(sid, res["items"]); q.recompute_subtotal(sid)
        import reports.excel_report as er
        er.OUTPUT_DIR = _util.FIX_DIR
        out = er.export_level_tree_xlsx(sid)
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    grid = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
    hdr = grid[0][:6]
    total_row = grid[-1]
    # 납블록 행: 대분류=재료비, 품명=납블록, 금액=600. 부품 행 2개(파트금액), 합계=600
    part_rows = [g for g in grid if str(g[5]).startswith("↳")]
    ok = (hdr == ["대분류", "중분류", "소분류", "세분류", "품명", "부품"]
          and len(part_rows) == 2
          and all(pr[10] == "" for pr in part_rows)   # 부품 행 금액칸 빈칸(합계 제외)
          and total_row[10] == 600)
    rec.add("14-3", "레벨별 엑셀 추출 (#3)",
            "대/중/소/세/품명/부품 열 + 부품 합계 제외",
            "헤더 6열, 부품행 2(금액 빈칸), 합계=600(품목만)",
            f"hdr={hdr} 부품행={len(part_rows)} 합계={total_row[10]}",
            "PASS" if ok else "FAIL",
            "" if ok else "엑셀 열/부품제외/합계 오류")

    return rec.flush()


if __name__ == "__main__":
    run()
